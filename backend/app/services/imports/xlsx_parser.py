from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.imports.errors import ImportFileError
from app.services.imports.limits import MAX_IMPORT_BYTES, MAX_IMPORT_ROWS
from app.services.imports.normalization import map_headers
from app.services.imports.parsed import ParsedImport, ParsedRow


def _ensure_size(content: bytes) -> None:
    if len(content) > MAX_IMPORT_BYTES:
        raise ImportFileError("El archivo supera el tamaño máximo permitido.")


def _cell(value: object) -> object:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def parse_xlsx(content: bytes) -> ParsedImport:
    if not content:
        raise ImportFileError("El archivo está vacío.")

    _ensure_size(content)

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ImportFileError("No se pudo leer el archivo Excel.") from exc

    try:
        if not workbook.worksheets:
            raise ImportFileError("El archivo Excel no contiene hojas.")

        worksheet = workbook.worksheets[0]
        rows_iterator = worksheet.iter_rows(values_only=True)
        header: tuple[object, ...] | None = None
        header_line = 0

        for row_number, row in enumerate(rows_iterator, start=1):
            values = tuple(_cell(value) for value in row)
            if any(value is not None for value in values):
                header = values
                header_line = row_number
                break

        if header is None:
            raise ImportFileError("El archivo Excel no contiene una cabecera válida.")

        mapping = map_headers(header)
        if not mapping.is_valid:
            raise ImportFileError(
                "El archivo no tiene las columnas requeridas.",
                mapping.errors,
            )

        rows: list[ParsedRow] = []
        for row_number, row in enumerate(rows_iterator, start=header_line + 1):
            values = tuple(_cell(value) for value in row)
            parsed = ParsedRow(
                row_number=row_number,
                fecha=_cell(values[mapping.positions["fecha"]])
                if mapping.positions["fecha"] < len(values)
                else None,
                descripcion=(
                    _cell(values[mapping.positions["descripcion"]])
                    if mapping.positions["descripcion"] < len(values)
                    else None
                ),
                valor=_cell(values[mapping.positions["valor"]])
                if mapping.positions["valor"] < len(values)
                else None,
                categoria=(
                    _cell(values[mapping.positions["categoria"]])
                    if "categoria" in mapping.positions
                    and mapping.positions["categoria"] < len(values)
                    else None
                ),
            )
            if (
                parsed.fecha is None
                and parsed.descripcion is None
                and parsed.valor is None
                and parsed.categoria is None
            ):
                continue
            rows.append(parsed)
            if len(rows) > MAX_IMPORT_ROWS:
                raise ImportFileError("El archivo supera el máximo de filas permitido.")

        return ParsedImport(headers=mapping.positions, rows=tuple(rows))
    finally:
        workbook.close()
